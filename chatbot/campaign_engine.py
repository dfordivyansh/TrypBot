# =========================================================
# FILE: chatbot/campaign_engine.py
# =========================================================
# WhatsApp broadcast campaigns — the non-view logic.
#
#   • phone normalisation + CSV / Excel contact-file parsing
#   • audience resolution (tags, labels, dates, segments …)
#   • recipient materialisation (variables resolved per contact)
#   • the sender (template messages via the WhatsApp Cloud API)
#   • webhook hooks  (delivery status, replies, STOP / SUBSCRIBE)
#   • the background worker (in-process thread or `run_campaigns`)
#
# Safe to run more than one worker: every recipient is claimed with an
# atomic UPDATE ... WHERE status='pending', so no message is sent twice.
# =========================================================

import csv
import io
import os
import re
import threading
import time
import traceback
from datetime import datetime, timedelta, timezone as dt_timezone
from decimal import Decimal, InvalidOperation

import requests
from django.conf import settings
from django.db import close_old_connections, transaction
from django.db.models import Count, F, Q
from django.utils import timezone
from django.utils.dateparse import parse_date

from .models import (
    Campaign,
    CampaignRecipient,
    ChatMessage,
    ChatSession,
    ContactTag,
    OptInRecord,
)


# =========================================================
# CONFIG
# =========================================================

GRAPH_VERSION        = "v22.0"
DEFAULT_COUNTRY_CODE = getattr(settings, "CAMPAIGN_DEFAULT_COUNTRY_CODE", "91")
SEND_INTERVAL        = float(getattr(settings, "CAMPAIGN_SEND_INTERVAL", 0.15))   # seconds between sends
MAX_RECIPIENTS       = int(getattr(settings, "CAMPAIGN_MAX_RECIPIENTS", 20000))
CSV_MAX_ROWS         = int(getattr(settings, "CAMPAIGN_CSV_MAX_ROWS", 5000))
CSV_MAX_BYTES        = 5 * 1024 * 1024
BATCH_PER_PASS       = 25      # recipients handled per campaign per worker pass (fairness)
MAX_ATTEMPTS         = 3
STALE_SENDING_MIN    = 10
THROTTLE_SLEEP       = 20      # seconds to back off when Meta rate-limits us
CONSECUTIVE_FAIL_MAX = 10

# Meta error codes ------------------------------------------------------
# Rate limiting / transient  → back off and try the same contact again
RETRY_CODES = {"130429", "131056", "80007", "4", "17", "613", "131000", "2", "1"}
THROTTLE_CODES = {"130429", "131056", "80007", "4", "17", "613"}
# Problems that affect *every* message → stop the campaign so the list isn't burned
FATAL_CODES = {
    "190", "102", "10",            # token invalid / expired / no permission
    "132000", "132001", "132005",  # template params / template missing / text too long
    "132007", "132012", "132015", "132016",
    "131008",                      # required parameter missing (e.g. header media)
    "131031", "131042", "131048",  # account locked / payment issue / spam rate limit
    "368",                         # temporarily blocked
}
# Per-contact problems that are normal in a stale list (excluded from the
# "N failures in a row" safety stop)
PER_RECIPIENT_CODES = {"131026", "131021", "131030", "131049", "131009", "131051", "100"}

OPT_OUT_KEYWORDS = {"stop", "stop all", "unsubscribe", "opt out", "optout",
                    "stop messages", "stop promotions", "stop messaging"}
OPT_IN_KEYWORDS  = {"subscribe", "resubscribe", "opt in", "optin", "start promotions"}

OPT_OUT_REPLY = (
    "You have been unsubscribed from promotional messages. "
    "Reply SUBSCRIBE anytime to start receiving them again."
)
OPT_IN_REPLY = "Welcome back! You will receive our offers and updates again."


def _log(msg):
    print(f"[CAMPAIGN] {msg}")


# =========================================================
# PHONE NUMBERS
# =========================================================

def normalize_phone(raw, default_cc=None):
    """
    Return a WhatsApp-style number (digits only, with country code) or None.

    Accepts +91 98765 43210, 09876543210, 9876543210 (→ default country
    code), 0091…, and Excel artefacts such as "919876543210.0" / 9.19876543210E+11.
    """
    default_cc = default_cc or DEFAULT_COUNTRY_CODE
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None

    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    elif re.fullmatch(r"\d+(\.\d+)?[eE][+-]?\d+", s):
        try:
            s = str(int(Decimal(s)))
        except (InvalidOperation, ValueError):
            return None

    had_plus = s.startswith("+")
    digits = re.sub(r"\D", "", s)
    if not digits:
        return None

    if digits.startswith("00"):
        digits = digits[2:]
        had_plus = True

    if not had_plus:
        if len(digits) == 10:
            digits = default_cc + digits
        elif len(digits) == 11 and digits.startswith("0"):
            digits = default_cc + digits[1:]

    if not (11 <= len(digits) <= 15) or digits.startswith("0"):
        return None

    # Indian mobile numbers start with 6-9
    if digits.startswith("91") and len(digits) == 12 and digits[2] not in "6789":
        return None

    return digits


def _phone_variants(phone):
    v = {phone, "+" + phone}
    if phone.startswith(DEFAULT_COUNTRY_CODE) and len(phone) == len(DEFAULT_COUNTRY_CODE) + 10:
        v.add(phone[len(DEFAULT_COUNTRY_CODE):])
    return v


def sessions_by_phone(phones):
    """{normalised_phone: ChatSession} for the numbers that already exist in the CRM."""
    phones = list(set(phones))
    variant_to_phone = {}
    for p in phones:
        for v in _phone_variants(p):
            variant_to_phone[v] = p

    wanted = set(phones)
    found = {}
    variants = list(variant_to_phone)
    for i in range(0, len(variants), 500):
        chunk = variants[i:i + 500]
        for s in ChatSession.objects.filter(phone_number__in=chunk).select_related("lead").order_by("id"):
            key = normalize_phone(s.phone_number)
            if key in wanted and key not in found:
                found[key] = s
    return found


# =========================================================
# CONTACT FILE PARSING  (CSV / XLSX)
# =========================================================

_PHONE_HEADERS = {"phone", "phonenumber", "mobile", "mobilenumber", "mobileno", "phoneno", "number",
                  "whatsapp", "whatsappnumber", "contact", "contactnumber", "contactno", "cell",
                  "telephone", "tel", "msisdn"}
_NAME_HEADERS  = {"name", "fullname", "customername", "contactname", "username", "firstname", "customer"}
_TAG_HEADERS   = {"tag", "tags", "label", "labels"}


def _hnorm(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _read_table(uploaded):
    """Return a list of rows (lists of str) from an uploaded .csv / .xlsx / .txt file."""
    name = (uploaded.name or "").lower()
    raw = uploaded.read()
    if len(raw) > CSV_MAX_BYTES:
        raise ValueError("File is larger than 5 MB.")

    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if c is None else (str(int(c)) if isinstance(c, float) and c.is_integer() else str(c)) for c in r]
            if any(x.strip() for x in cells):
                rows.append(cells)
        return rows

    if name.endswith(".xls"):
        raise ValueError("Old .xls files are not supported — save the sheet as .xlsx or .csv.")

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [r for r in reader if any((c or "").strip() for c in r)]


def parse_contacts_file(uploaded):
    """
    Parse an uploaded contact list.

    Returns {"columns": [...], "rows": [{"phone","name","tags","extra"}],
             "invalid": [{"row","value","reason"}], "duplicates": int, "total_rows": int}
    Raises ValueError with a user-friendly message for unusable files.
    """
    table = _read_table(uploaded)
    if not table:
        raise ValueError("The file is empty.")

    first = table[0]
    hn = [_hnorm(c) for c in first]
    phone_idx = next((i for i, h in enumerate(hn) if h in _PHONE_HEADERS), None)
    name_idx = next((i for i, h in enumerate(hn) if h in _NAME_HEADERS), None)
    tag_idx = next((i for i, h in enumerate(hn) if h in _TAG_HEADERS), None)

    has_header = True
    if phone_idx is None:
        # No header row: is the first cell of row 1 itself a phone number?
        if normalize_phone(first[0]) is not None:
            has_header = False
            headers = [f"Column {i + 1}" for i in range(len(first))]
            data = table
            phone_idx, name_idx = 0, (1 if len(first) > 1 else None)
        else:
            raise ValueError(
                "Could not find a phone column. Name the column header “phone” "
                "(or mobile / whatsapp / number)."
            )
    else:
        headers = [str(c).strip() or f"Column {i + 1}" for i, c in enumerate(first)]
        data = table[1:]

    if len(data) > CSV_MAX_ROWS:
        raise ValueError(f"The file has {len(data)} rows — the limit is {CSV_MAX_ROWS} per upload. Please split it.")

    rows, invalid, seen, dups = [], [], set(), 0
    for n, r in enumerate(data, start=2 if has_header else 1):
        cell = r[phone_idx] if phone_idx < len(r) else ""
        phone = normalize_phone(cell)
        if not phone:
            invalid.append({"row": n, "value": str(cell)[:40], "reason": "Invalid phone number"})
            continue
        if phone in seen:
            dups += 1
            continue
        seen.add(phone)

        name = (r[name_idx].strip() if name_idx is not None and name_idx < len(r) else "")[:255]
        tags = []
        if tag_idx is not None and tag_idx < len(r):
            tags = [t.strip() for t in re.split(r"[,;|]", r[tag_idx]) if t.strip()][:10]
        extra = {}
        for i, h in enumerate(headers):
            if i in (phone_idx,) or i >= len(r):
                continue
            v = str(r[i]).strip()
            if v:
                extra[h] = v[:500]
        rows.append({"phone": phone, "name": name, "tags": tags, "extra": extra})

    return {
        "columns": [h for i, h in enumerate(headers) if i != phone_idx],
        "rows": rows,
        "invalid": invalid,
        "duplicates": dups,
        "total_rows": len(data),
    }


# =========================================================
# TAGS / OPT-OUT HELPERS
# =========================================================

def get_or_create_tag(name, user=None, color="#25D366"):
    name = (name or "").strip()[:60]
    if not name:
        return None
    tag = ContactTag.objects.filter(name__iexact=name).first()
    if tag:
        return tag
    return ContactTag.objects.create(name=name, color=color, created_by=user)


def set_opt_out(sessions_qs, opted_out):
    """Bulk opt-out / opt-in. Returns rows changed."""
    now = timezone.now() if opted_out else None
    return ChatSession.objects.filter(pk__in=sessions_qs.values("pk")).update(
        marketing_opt_out=bool(opted_out), marketing_opt_out_at=now)


# =========================================================
# TEMPLATE VARIABLES
# =========================================================

_PLACEHOLDER_RE = re.compile(r"\{\{\s*(\d+)\s*\}\}")

VARIABLE_SOURCES = ("name", "first_name", "phone", "destination", "email", "column", "fixed")

DEFAULT_FALLBACKS = {
    "name": "there", "first_name": "there", "phone": "-", "destination": "your destination",
    "email": "-", "column": "-", "fixed": "-",
}


def template_placeholders(text):
    """Sorted unique placeholder numbers used in a template body: '{{1}} … {{2}}' → [1, 2]."""
    return sorted({int(n) for n in _PLACEHOLDER_RE.findall(text or "")})


def render_template_text(text, variables):
    return _PLACEHOLDER_RE.sub(lambda m: str(variables.get(m.group(1), m.group(0))), text or "")


def clean_param(value):
    """WhatsApp rejects newlines/tabs and long runs of spaces inside template parameters."""
    s = re.sub(r"[\r\n\t]+", " ", str(value or ""))
    s = re.sub(r" {4,}", "   ", s).strip()
    return s[:1000]


def recipient_context(session=None, phone="", name="", extra=None):
    ctx = {"name": name or "", "phone": phone or "", "destination": "", "email": "", "extra": extra or {}}
    if session is not None:
        ctx["name"] = (session.user_name or name or "").strip()
        ctx["phone"] = session.phone_number or phone
        lead = getattr(session, "lead", None)
        ctx["destination"] = (getattr(lead, "destination", "") or session.current_destination or "").strip()
        ctx["email"] = (getattr(lead, "email", "") or "").strip()
    return ctx


def compute_variables(mapping, placeholders, ctx):
    """{"1": "Rahul", "2": "Goa"} — every value is guaranteed non-empty."""
    out = {}
    for n in placeholders:
        spec = (mapping or {}).get(str(n), {}) or {}
        src = spec.get("source", "name")
        val = ""
        if src == "name":
            val = ctx["name"]
        elif src == "first_name":
            val = ctx["name"].split()[0] if ctx["name"].split() else ""
        elif src == "phone":
            val = ctx["phone"]
        elif src == "destination":
            val = ctx["destination"]
        elif src == "email":
            val = ctx["email"]
        elif src == "column":
            val = ctx["extra"].get(spec.get("value", ""), "")
        elif src == "fixed":
            val = spec.get("value", "")
        val = clean_param(val)
        if not val:
            val = clean_param(spec.get("fallback", "")) or DEFAULT_FALLBACKS.get(src, "-")
        out[str(n)] = val
    return out


def validate_mapping(mapping, placeholders):
    """Return an error string or None."""
    for n in placeholders:
        spec = (mapping or {}).get(str(n))
        if not spec:
            return f"Variable {{{{{n}}}}} has no value assigned."
        src = spec.get("source")
        if src not in VARIABLE_SOURCES:
            return f"Variable {{{{{n}}}}} has an invalid source."
        if src in ("column", "fixed") and not str(spec.get("value", "")).strip():
            return f"Variable {{{{{n}}}}}: please enter the {'column name' if src == 'column' else 'text'}."
    return None


# =========================================================
# AUDIENCE
# =========================================================

def _int_list(v):
    out = []
    for x in v or []:
        try:
            out.append(int(x))
        except (TypeError, ValueError):
            pass
    return out


def base_contacts_qs(user, role=None):
    """Contacts the user is allowed to message (sales → own chats only)."""
    from .crm_views import get_user_role

    role = role or get_user_role(user)
    qs = ChatSession.objects.exclude(phone_number__isnull=True).exclude(phone_number="")
    if role == "sales":
        qs = qs.filter(assigned_sales_person=user)
    return qs


def resolve_audience(filters, user, role=None):
    """ChatSession queryset for a filter dict (see Segment.filters)."""
    from .crm_views import get_user_role

    role = role or get_user_role(user)
    qs = base_contacts_qs(user, role)
    f = filters or {}

    q = (f.get("q") or "").strip()
    if q:
        qs = qs.filter(Q(user_name__icontains=q) | Q(phone_number__icontains=q))

    tags = _int_list(f.get("tags"))
    if tags:
        if f.get("tag_mode") == "all":
            for t in tags:
                qs = qs.filter(tags__id=t)
        else:
            qs = qs.filter(tags__id__in=tags)

    ex = _int_list(f.get("exclude_tags"))
    if ex:
        qs = qs.exclude(tags__id__in=ex)

    labels = _int_list(f.get("labels"))
    if labels:
        qs = qs.filter(conv_status__label_id__in=labels)

    if role != "sales":
        assigned = _int_list(f.get("assigned_to"))
        if assigned:
            qs = qs.filter(assigned_sales_person_id__in=assigned)

    d_from = parse_date(f.get("created_from") or "") if f.get("created_from") else None
    d_to = parse_date(f.get("created_to") or "") if f.get("created_to") else None
    if d_from:
        qs = qs.filter(created_at__date__gte=d_from)
    if d_to:
        qs = qs.filter(created_at__date__lte=d_to)

    dest = (f.get("destination") or "").strip()
    if dest:
        qs = qs.filter(Q(lead__destination__icontains=dest) | Q(current_destination__icontains=dest))

    source = (f.get("source") or "").strip()
    if source:
        qs = qs.filter(lead__source__iexact=source)

    try:
        within = int(f.get("active_within_days") or 0)
    except (TypeError, ValueError):
        within = 0
    if within > 0:
        qs = qs.filter(last_customer_message_time__gte=timezone.now() - timedelta(days=within))

    try:
        inactive = int(f.get("inactive_days") or 0)
    except (TypeError, ValueError):
        inactive = 0
    if inactive > 0:
        cut = timezone.now() - timedelta(days=inactive)
        qs = qs.filter(Q(last_customer_message_time__lt=cut) | Q(last_customer_message_time__isnull=True))

    return qs.distinct()


def audience_preview(qs, sample=5):
    total = qs.count()
    opted_out = qs.filter(marketing_opt_out=True).count()
    eligible_qs = qs.filter(marketing_opt_out=False).select_related("lead").order_by("-updated_at")
    sample_rows = []
    for s in eligible_qs[:sample]:
        ctx = recipient_context(s)
        sample_rows.append({"name": ctx["name"], "phone": ctx["phone"],
                            "destination": ctx["destination"], "email": ctx["email"]})
    return {"total": total, "opted_out": opted_out, "eligible": total - opted_out, "sample": sample_rows}


# =========================================================
# RECIPIENT MATERIALISATION
# =========================================================

class AudienceError(ValueError):
    pass


def _entries_from_sessions(sessions):
    """Dedupe sessions by normalised phone → list of entry dicts."""
    seen, entries = set(), []
    for s in sessions:
        phone = normalize_phone(s.phone_number)
        if not phone or phone in seen:
            continue
        seen.add(phone)
        entries.append({"phone": phone, "session": s, "name": s.user_name or "", "extra": {}})
    return entries


def _entries_from_csv(rows, user, role, tag_names, create_contacts=True, consent_text=""):
    """
    Match CSV rows to existing contacts and (optionally) create the missing
    ones so campaign replies land in the normal inbox.
    """
    clean, seen = [], set()
    for r in rows:
        phone = normalize_phone(r.get("phone"))
        if not phone or phone in seen:
            continue
        seen.add(phone)
        clean.append({"phone": phone, "name": (r.get("name") or "")[:255],
                      "extra": r.get("extra") or {}, "tags": r.get("tags") or []})

    existing = sessions_by_phone([c["phone"] for c in clean])
    tag_cache = {}

    def tag_for(n):
        k = (n or "").strip().lower()
        if k not in tag_cache:
            tag_cache[k] = get_or_create_tag(n, user)
        return tag_cache[k]

    tag_objs = [t for t in (tag_for(n) for n in tag_names or []) if t]

    entries = []
    for c in clean:
        session = existing.get(c["phone"])
        if session is None and create_contacts:
            session, _ = ChatSession.objects.get_or_create(
                session_id=c["phone"],
                defaults={
                    "phone_number": c["phone"],
                    "user_name": c["name"] or None,
                    "assigned_sales_person": user if role == "sales" else None,
                },
            )
            if not session.phone_number:
                session.phone_number = c["phone"]
                session.save(update_fields=["phone_number"])
            OptInRecord.objects.create(
                phone_number=c["phone"], opt_in_method="import",
                session=session, consent_text=consent_text[:500],
            )
        elif session is not None and not session.user_name and c["name"]:
            ChatSession.objects.filter(pk=session.pk).update(user_name=c["name"])
            session.user_name = c["name"]

        if session is not None:
            row_tags = list(tag_objs) + [t for t in (tag_for(n) for n in c["tags"]) if t]
            if row_tags:
                session.tags.add(*row_tags)

        entries.append({"phone": c["phone"], "session": session, "name": c["name"], "extra": c["extra"]})
    return entries


@transaction.atomic
def import_contacts(rows, user, role, tag_names=None, consent_text=""):
    """Add uploaded contacts to the CRM (no message is sent). Returns {"created": n, "existing": m}."""
    phones = [normalize_phone(r.get("phone")) for r in rows]
    pre = sessions_by_phone([p for p in phones if p])
    entries = _entries_from_csv(rows, user, role, tag_names=tag_names or [],
                                create_contacts=True, consent_text=consent_text)
    created = sum(1 for e in entries if e["phone"] not in pre)
    return {"created": created, "existing": len(entries) - created}


@transaction.atomic
def create_campaign(user, role, data):
    """
    Create a campaign and materialise its recipients.

    `data` is the JSON payload from the wizard. Returns (campaign, summary).
    Raises AudienceError / ValueError with user-facing messages.
    """
    from .models import CRMTemplate, Segment

    name = (data.get("name") or "").strip()
    if not name:
        raise ValueError("Please give the campaign a name.")

    template = CRMTemplate.objects.filter(pk=data.get("template_id")).first()
    if not template:
        raise ValueError("Please choose a template.")
    if not template.template_name:
        raise ValueError("This template has no WhatsApp template name. Add the approved name from Meta first.")

    placeholders = template_placeholders(template.template_text)
    mapping = data.get("mapping") or {}
    err = validate_mapping(mapping, placeholders)
    if err:
        raise ValueError(err)

    header = data.get("header") or {}
    h_type = header.get("type") or ""
    h_url = (header.get("url") or "").strip()
    if h_type not in ("", "image", "video", "document"):
        raise ValueError("Invalid header media type.")
    if h_type and not h_url.lower().startswith("https://"):
        raise ValueError("Header media needs a public https:// link.")

    schedule = data.get("schedule") or {}
    mode = schedule.get("mode", "now")
    scheduled_at = None
    if mode == "later":
        try:
            scheduled_at = datetime.fromisoformat(str(schedule.get("at")).replace("Z", "+00:00"))
        except ValueError:
            raise ValueError("Please choose a valid schedule date and time.")
        if timezone.is_naive(scheduled_at):
            scheduled_at = timezone.make_aware(scheduled_at)
        if scheduled_at <= timezone.now() + timedelta(seconds=30):
            raise ValueError("The scheduled time must be in the future.")
    elif mode not in ("now", "draft"):
        raise ValueError("Invalid schedule option.")

    if mode != "draft" and not data.get("consent"):
        raise ValueError("Please confirm that these contacts have opted in to receive WhatsApp messages.")

    aud = data.get("audience") or {}
    a_type = aud.get("type", "filters")
    segment = None
    filters = {}

    if a_type == "csv":
        entries = _entries_from_csv(
            aud.get("csv_rows") or [], user, role,
            tag_names=aud.get("csv_tags") or [],
            create_contacts=aud.get("create_contacts", True),
            consent_text=f"Import confirmed by {user.username} for campaign '{name}'",
        )
    else:
        if a_type == "segment":
            segment = Segment.objects.filter(pk=aud.get("segment_id")).first()
            if not segment:
                raise ValueError("Segment not found.")
            filters = segment.filters
        else:
            filters = aud.get("filters") or {}
        qs = resolve_audience(filters, user, role)
        ids = _int_list(aud.get("session_ids"))
        if a_type == "contacts":
            if not ids:
                raise ValueError("Select at least one contact.")
            qs = qs.filter(id__in=ids)
        entries = _entries_from_sessions(qs.select_related("lead"))

    if not entries:
        raise AudienceError("No contacts matched this audience.")
    if len(entries) > MAX_RECIPIENTS:
        raise AudienceError(f"Too many recipients ({len(entries)}). The limit is {MAX_RECIPIENTS} per campaign.")

    campaign = Campaign.objects.create(
        name=name[:255],
        template=template,
        template_title=template.title,
        template_name=template.template_name,
        template_language=template.language or "en",
        template_text=template.template_text,
        variable_mapping=mapping,
        header_media_type=h_type,
        header_media_url=h_url,
        header_media_filename=(header.get("filename") or "")[:255],
        audience_type=a_type,
        audience_filters=filters if a_type != "csv" else {},
        segment=segment,
        status="draft",
        scheduled_at=scheduled_at,
        consent_confirmed=bool(data.get("consent")),
        created_by=user,
    )

    recips, opted_out = [], 0
    for e in entries:
        s = e["session"]
        ctx = recipient_context(s, e["phone"], e["name"], e["extra"])
        variables = compute_variables(mapping, placeholders, ctx)
        skip = s is not None and s.marketing_opt_out
        opted_out += 1 if skip else 0
        recips.append(CampaignRecipient(
            campaign=campaign, session=s, phone=e["phone"], name=ctx["name"][:255],
            variables=variables, extra=e["extra"],
            status="skipped" if skip else "pending",
            error_message="Contact has opted out of marketing messages" if skip else "",
        ))
    CampaignRecipient.objects.bulk_create(recips, batch_size=500)

    eligible = len(recips) - opted_out
    if eligible == 0 and mode != "draft":
        raise AudienceError("Every contact in this audience has opted out.")

    if mode != "draft":
        launch_campaign(campaign, scheduled_at)

    return campaign, {"recipients": len(recips), "eligible": eligible, "opted_out": opted_out}


# =========================================================
# STATE TRANSITIONS
# =========================================================

def launch_campaign(campaign, scheduled_at=None):
    now = timezone.now()
    if scheduled_at and scheduled_at > now:
        campaign.status = "scheduled"
        campaign.scheduled_at = scheduled_at
        campaign.save(update_fields=["status", "scheduled_at", "updated_at"])
    else:
        campaign.status = "running"
        campaign.scheduled_at = None
        campaign.started_at = campaign.started_at or now
        campaign.last_error = ""
        campaign.save(update_fields=["status", "scheduled_at", "started_at", "last_error", "updated_at"])


def pause_campaign(campaign):
    if campaign.status not in ("running", "scheduled"):
        raise ValueError("Only running or scheduled campaigns can be paused.")
    Campaign.objects.filter(pk=campaign.pk, status__in=("running", "scheduled")).update(status="paused", updated_at=timezone.now())


def resume_campaign(campaign):
    if campaign.status not in ("paused", "failed"):
        raise ValueError("Only paused or failed campaigns can be resumed.")
    if not campaign.recipients.filter(status__in=("pending", "sending")).exists():
        raise ValueError("Nothing left to send. Use “Retry failed” to resend failed messages.")
    if campaign.scheduled_at and campaign.scheduled_at > timezone.now():
        launch_campaign(campaign, campaign.scheduled_at)
    else:
        launch_campaign(campaign)


def cancel_campaign(campaign):
    if campaign.status in ("completed", "cancelled"):
        raise ValueError("This campaign has already finished.")
    campaign.recipients.filter(status__in=("pending", "sending")).update(
        status="skipped", error_message="Campaign cancelled")
    Campaign.objects.filter(pk=campaign.pk).update(
        status="cancelled", completed_at=timezone.now(), updated_at=timezone.now())


def retry_failed(campaign):
    if campaign.status not in ("completed", "paused", "failed"):
        raise ValueError("Retry is available once the campaign has finished or is paused.")
    n = campaign.recipients.filter(status="failed").update(
        status="pending", attempts=0, error_code="", error_message="", claimed_at=None)
    if n == 0:
        raise ValueError("There are no failed messages to retry.")
    campaign.completed_at = None
    campaign.save(update_fields=["completed_at", "updated_at"])
    launch_campaign(campaign)
    return n


def duplicate_campaign(campaign, user):
    """A draft copy with all recipients reset to pending (opted-out contacts stay skipped)."""
    with transaction.atomic():
        copy = Campaign.objects.create(
            name=f"{campaign.name} (copy)"[:255],
            template=campaign.template, template_title=campaign.template_title,
            template_name=campaign.template_name, template_language=campaign.template_language,
            template_text=campaign.template_text, variable_mapping=campaign.variable_mapping,
            header_media_type=campaign.header_media_type, header_media_url=campaign.header_media_url,
            header_media_filename=campaign.header_media_filename,
            audience_type=campaign.audience_type, audience_filters=campaign.audience_filters,
            segment=campaign.segment, status="draft", created_by=user,
        )
        rows = []
        for r in campaign.recipients.select_related("session").all().iterator():
            skip = r.session is not None and r.session.marketing_opt_out
            rows.append(CampaignRecipient(
                campaign=copy, session=r.session, phone=r.phone, name=r.name,
                variables=r.variables, extra=r.extra,
                status="skipped" if skip else "pending",
                error_message="Contact has opted out of marketing messages" if skip else "",
            ))
        CampaignRecipient.objects.bulk_create(rows, batch_size=500)
    return copy


# =========================================================
# STATS
# =========================================================

def stats_annotations(prefix="recipients__"):
    p = prefix
    return {
        "r_total":     Count(f"{p}id"),
        "r_pending":   Count(f"{p}id", filter=Q(**{f"{p}status__in": ("pending", "sending")})),
        "r_sent":      Count(f"{p}id", filter=Q(**{f"{p}sent_at__isnull": False})),
        "r_delivered": Count(f"{p}id", filter=Q(**{f"{p}delivered_at__isnull": False})),
        "r_read":      Count(f"{p}id", filter=Q(**{f"{p}read_at__isnull": False})),
        "r_replied":   Count(f"{p}id", filter=Q(**{f"{p}replied_at__isnull": False})),
        "r_failed":    Count(f"{p}id", filter=Q(**{f"{p}status": "failed"})),
        "r_skipped":   Count(f"{p}id", filter=Q(**{f"{p}status": "skipped"})),
    }


def campaign_stats(campaign):
    agg = campaign.recipients.aggregate(**stats_annotations(prefix=""))
    return stats_dict(agg)


def stats_dict(agg):
    total = agg["r_total"] or 0
    skipped = agg["r_skipped"] or 0
    sendable = max(total - skipped, 0)
    done = max(sendable - (agg["r_pending"] or 0), 0)
    return {
        "total": total, "sendable": sendable, "pending": agg["r_pending"] or 0,
        "sent": agg["r_sent"] or 0, "delivered": agg["r_delivered"] or 0,
        "read": agg["r_read"] or 0, "replied": agg["r_replied"] or 0,
        "failed": agg["r_failed"] or 0, "skipped": skipped,
        "processed": done,
        "progress": round(done * 100 / sendable) if sendable else 100,
    }


# =========================================================
# SENDER
# =========================================================

def build_components(campaign, variables):
    comps = []
    if campaign.header_media_type and campaign.header_media_url:
        kind = campaign.header_media_type
        media = {"link": campaign.header_media_url}
        if kind == "document" and campaign.header_media_filename:
            media["filename"] = campaign.header_media_filename
        comps.append({"type": "header", "parameters": [{"type": kind, kind: media}]})

    nums = template_placeholders(campaign.template_text)
    if nums:
        comps.append({
            "type": "body",
            "parameters": [{"type": "text", "text": variables.get(str(n)) or "-"} for n in nums],
        })
    return comps


def graph_send_template(to, name, language, components):
    """
    POST one template message. Returns a dict:
      ok, id | code, message, retry, throttle, fatal
    """
    url = f"https://graph.facebook.com/{GRAPH_VERSION}/{settings.WHATSAPP_PHONE_NUMBER_ID}/messages"
    headers = {
        "Authorization": f"Bearer {settings.WHATSAPP_ACCESS_TOKEN}",
        "Content-Type": "application/json",
    }
    payload = {
        "messaging_product": "whatsapp",
        "to": str(to),
        "type": "template",
        "template": {"name": name, "language": {"code": language or "en"}},
    }
    if components:
        payload["template"]["components"] = components

    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=30)
    except requests.RequestException as exc:
        return {"ok": False, "code": "network", "message": f"Network error: {str(exc)[:200]}",
                "retry": True, "throttle": False, "fatal": False}

    try:
        data = resp.json()
    except ValueError:
        data = {}

    if resp.ok and data.get("messages"):
        return {"ok": True, "id": data["messages"][0].get("id", "")}

    err = data.get("error", {}) or {}
    code = str(err.get("code", resp.status_code))
    msg = err.get("message") or f"HTTP {resp.status_code}"
    details = (err.get("error_data") or {}).get("details")
    if details and details not in msg:
        msg = f"{msg} — {details}"
    return {
        "ok": False, "code": code, "message": msg[:500],
        "retry": code in RETRY_CODES or resp.status_code in (429, 500, 502, 503, 504),
        "throttle": code in THROTTLE_CODES or resp.status_code == 429,
        "fatal": code in FATAL_CODES or resp.status_code == 401,
    }


_consecutive = {}     # campaign_id -> (code, count)   (per-process safety net)


def _fail_campaign(campaign_id, reason):
    Campaign.objects.filter(pk=campaign_id, status="running").update(
        status="failed", last_error=reason[:1000], updated_at=timezone.now())
    _log(f"campaign {campaign_id} stopped: {reason}")


def _send_one(campaign, rec):
    """Send to one claimed recipient. Returns 'ok' | 'fatal' | 'throttled'."""
    session = rec.session
    if session is not None and ChatSession.objects.filter(pk=session.pk, marketing_opt_out=True).exists():
        CampaignRecipient.objects.filter(pk=rec.pk).update(
            status="skipped", error_message="Contact opted out before this message was sent")
        return "ok"

    if not settings.WHATSAPP_PHONE_NUMBER_ID or not settings.WHATSAPP_ACCESS_TOKEN:
        CampaignRecipient.objects.filter(pk=rec.pk).update(status="pending", attempts=F("attempts") - 1)
        _fail_campaign(campaign.pk, "WhatsApp credentials (WHATSAPP_ACCESS_TOKEN / WHATSAPP_PHONE_NUMBER_ID) are not configured.")
        return "fatal"

    components = build_components(campaign, rec.variables or {})
    result = graph_send_template(rec.phone, campaign.template_name, campaign.template_language, components)

    if result["ok"]:
        now = timezone.now()
        CampaignRecipient.objects.filter(pk=rec.pk).update(
            status="sent", whatsapp_message_id=result["id"], sent_at=now,
            error_code="", error_message="")
        _consecutive.pop(campaign.pk, None)
        _record_chat_message(campaign, rec, result["id"])
        return "ok"

    code, message = result["code"], result["message"]

    if result["fatal"]:
        CampaignRecipient.objects.filter(pk=rec.pk).update(status="pending", attempts=F("attempts") - 1, claimed_at=None)
        _fail_campaign(campaign.pk, f"WhatsApp rejected the campaign (error {code}): {message}")
        return "fatal"

    if result["retry"] and rec.attempts < MAX_ATTEMPTS:
        CampaignRecipient.objects.filter(pk=rec.pk).update(status="pending", claimed_at=None)
        return "throttled" if (result["throttle"] or code == "network") else "ok"

    CampaignRecipient.objects.filter(pk=rec.pk).update(
        status="failed", error_code=code[:20], error_message=message)

    if code not in PER_RECIPIENT_CODES:
        last_code, count = _consecutive.get(campaign.pk, (None, 0))
        count = count + 1 if last_code == code else 1
        _consecutive[campaign.pk] = (code, count)
        if count >= CONSECUTIVE_FAIL_MAX:
            _consecutive.pop(campaign.pk, None)
            _fail_campaign(
                campaign.pk,
                f"Stopped after {CONSECUTIVE_FAIL_MAX} failures in a row (error {code}): {message} "
                "Fix the problem, then press Resume.")
            return "fatal"
    return "ok"


def _record_chat_message(campaign, rec, wamid):
    """Show the broadcast in the contact's chat, without bumping it to the top of the inbox."""
    if rec.session_id is None:
        return
    try:
        text = render_template_text(campaign.template_text, rec.variables or {})
        media_type = campaign.header_media_type
        ChatMessage.objects.create(
            session_id=rec.session_id,
            sender_type="sales",
            message=text,
            whatsapp_message_id=wamid,
            message_status="sent",
            media_url=campaign.header_media_url if media_type else "",
            media_type=media_type or "",
            media_filename=campaign.header_media_filename if media_type == "document" else "",
            raw_payload={"campaign_id": campaign.pk, "template": campaign.template_name},
        )
        ChatSession.objects.filter(pk=rec.session_id).update(last_message=text[:1000])
    except Exception:
        traceback.print_exc()


def process_campaign(campaign_id, batch=BATCH_PER_PASS):
    """Send up to `batch` pending messages for one running campaign. Returns messages attempted."""
    campaign = Campaign.objects.filter(pk=campaign_id).first()
    if campaign is None or campaign.status != "running":
        return 0

    ids = list(campaign.recipients.filter(status="pending").order_by("id").values_list("id", flat=True)[:batch])
    if not ids:
        _finish_if_done(campaign_id)
        return 0

    attempted = 0
    for rid in ids:
        current = Campaign.objects.filter(pk=campaign_id).values_list("status", flat=True).first()
        if current != "running":
            break

        claimed = CampaignRecipient.objects.filter(pk=rid, status="pending").update(
            status="sending", attempts=F("attempts") + 1, claimed_at=timezone.now())
        if not claimed:
            continue

        rec = CampaignRecipient.objects.select_related("session").get(pk=rid)
        outcome = _send_one(campaign, rec)
        attempted += 1

        if outcome == "fatal":
            break
        if outcome == "throttled":
            _log(f"campaign {campaign_id}: rate limited by WhatsApp, backing off {THROTTLE_SLEEP}s")
            time.sleep(THROTTLE_SLEEP)
            break
        time.sleep(SEND_INTERVAL)

    _finish_if_done(campaign_id)
    return attempted


def _finish_if_done(campaign_id):
    if CampaignRecipient.objects.filter(campaign_id=campaign_id, status__in=("pending", "sending")).exists():
        return
    Campaign.objects.filter(pk=campaign_id, status="running").update(
        status="completed", completed_at=timezone.now(), updated_at=timezone.now())


def _promote_due_campaigns():
    now = timezone.now()
    for c in Campaign.objects.filter(status="scheduled", scheduled_at__lte=now):
        Campaign.objects.filter(pk=c.pk, status="scheduled").update(
            status="running", started_at=c.started_at or now, updated_at=now)
        _log(f"campaign {c.pk} '{c.name}' is due — starting")


def _recover_stale():
    cutoff = timezone.now() - timedelta(minutes=STALE_SENDING_MIN)
    CampaignRecipient.objects.filter(status="sending", claimed_at__lt=cutoff).update(
        status="failed", error_code="interrupted",
        error_message="Sending was interrupted (server restart?). It may or may not have been delivered — use Retry failed to resend.")


def run_worker_once():
    """One scheduler pass. Returns the number of messages attempted."""
    _promote_due_campaigns()
    _recover_stale()
    attempted = 0
    for cid in list(Campaign.objects.filter(status="running").order_by("id").values_list("id", flat=True)):
        attempted += process_campaign(cid)
    return attempted


def worker_loop(stop_event=None, idle_sleep=3.0):
    _log("worker started")
    while not (stop_event is not None and stop_event.is_set()):
        attempted = 0
        try:
            close_old_connections()
            attempted = run_worker_once()
        except Exception:
            traceback.print_exc()
        finally:
            close_old_connections()
        if not attempted:
            time.sleep(idle_sleep)


_worker_lock = threading.Lock()
_worker_thread = None


def start_background_worker():
    """Start the in-process worker thread (idempotent). Disable with CAMPAIGN_INLINE_WORKER=0."""
    global _worker_thread
    if os.environ.get("CAMPAIGN_INLINE_WORKER", "1") == "0":
        return False
    with _worker_lock:
        if _worker_thread is not None and _worker_thread.is_alive():
            return False
        _worker_thread = threading.Thread(target=worker_loop, name="campaign-worker", daemon=True)
        _worker_thread.start()
        return True


# =========================================================
# WEBHOOK HOOKS
# =========================================================

_STATUS_RANK = {"sent": 1, "delivered": 2, "read": 3}


def on_status_update(message_id, status, timestamp=None, errors=None):
    """Called from the WhatsApp webhook for every delivery receipt. True if it belonged to a campaign."""
    if not message_id:
        return False
    rec = CampaignRecipient.objects.filter(whatsapp_message_id=message_id).first()
    if rec is None:
        return False

    try:
        when = datetime.fromtimestamp(int(timestamp), tz=dt_timezone.utc) if timestamp else timezone.now()
    except (TypeError, ValueError, OSError):
        when = timezone.now()

    updates = {}
    cur = _STATUS_RANK.get(rec.status, 0)

    if status == "failed":
        if rec.delivered_at is None and rec.status != "read":
            err = (errors or [{}])[0] or {}
            msg = err.get("message") or err.get("title") or "Delivery failed"
            details = (err.get("error_data") or {}).get("details")
            if details and details not in msg:
                msg = f"{msg} — {details}"
            updates.update(status="failed", error_code=str(err.get("code", ""))[:20], error_message=msg[:500])
    elif status == "delivered":
        if rec.delivered_at is None:
            updates["delivered_at"] = when
        if cur < 2:
            updates.update(status="delivered", error_code="", error_message="")
    elif status == "read":
        if rec.read_at is None:
            updates["read_at"] = when
        if rec.delivered_at is None:
            updates["delivered_at"] = when
        updates.update(status="read", error_code="", error_message="")
    elif status == "sent":
        if rec.sent_at is None:
            updates["sent_at"] = when
        if cur < 1 and rec.status in ("pending", "sending"):
            updates["status"] = "sent"

    if updates:
        CampaignRecipient.objects.filter(pk=rec.pk).update(**updates)
    return True


def on_inbound_message(session, text):
    """
    Called from the webhook after a customer message is saved.
    Attributes replies to the latest campaign and handles STOP / SUBSCRIBE.
    Returns "opted_out", "opted_in" or None.
    """
    result = None
    key = re.sub(r"[^a-z ]", "", (text or "").lower()).strip()
    key = re.sub(r"\s+", " ", key)

    if key in OPT_OUT_KEYWORDS:
        if not session.marketing_opt_out:
            ChatSession.objects.filter(pk=session.pk).update(
                marketing_opt_out=True, marketing_opt_out_at=timezone.now())
            session.marketing_opt_out = True
            try:
                from .spam_detection import SpamDetector
                SpamDetector.record_opt_out(session.phone_number or "", reason="Replied STOP")
            except Exception:
                pass
        result = "opted_out"
    elif key in OPT_IN_KEYWORDS:
        if session.marketing_opt_out:
            ChatSession.objects.filter(pk=session.pk).update(marketing_opt_out=False, marketing_opt_out_at=None)
            session.marketing_opt_out = False
        result = "opted_in"

    latest = (CampaignRecipient.objects
              .filter(session=session, replied_at__isnull=True,
                      sent_at__gte=timezone.now() - timedelta(days=7))
              .order_by("-sent_at").first())
    if latest:
        CampaignRecipient.objects.filter(pk=latest.pk).update(replied_at=timezone.now())

    return result
